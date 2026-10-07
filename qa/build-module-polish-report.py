"""Generate the 1.5.2 module ledger and document; no production writes."""
from pathlib import Path
import json, hashlib, argparse
from collections import Counter
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
ROOT = Path(__file__).resolve().parents[1]
ART = ROOT / 'artifacts/module-polish-20261001'
parser = argparse.ArgumentParser()
parser.add_argument('--final', action='store_true')
args = parser.parse_args()
OUT = Path('/Users/fengjinwen/Desktop/a股/细节打磨版1.5.2/模块打磨与独立审计报告.docx') if args.final else ART / 'word-draft/模块打磨与独立审计报告草稿.docx'
def read(name):
    path = ART / name
    return json.loads(path.read_text()) if path.exists() else None
fe = read('frontend-inventory.json')['modules']
be = read('backend-inventory.json')['reviewed_modules']
rt = read('root-inventory.json')['modules']
dep = read('dependency-implementation.json')
impl = {x:read(x+'-implementation.json') for x in ['frontend','backend','dependency','root']}
validation = read('test-validation.json')
audit = read('independent-lifecycle-audit.json')
native = read('delivery-native.json')
install = read('delivery-install.json')
if args.final and (any(x is None for x in [validation, audit, native, install]) or validation.get('status') != 'passed' or audit.get('verdict', {}).get('codeAndNativeApproved') is not True or native.get('status') != 'passed'):
    raise SystemExit('Final evidence missing; retain draft until independent audit, tests and delivery records exist.')
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin, section.bottom_margin = Inches(.7), Inches(.65)
section.left_margin = section.right_margin = Inches(.78)
section.footer_distance = Inches(.3)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3']:
    st = doc.styles[name]
    st.font.name = 'Arial Unicode MS'
    st.font.color.rgb = RGBColor(0, 0, 0)
    fonts = st.element.get_or_add_rPr().rFonts
    for attr in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme']:
        fonts.attrib.pop(qn('w:' + attr), None)
    fonts.set(qn('w:eastAsia'), 'Arial Unicode MS')
    for border in list(st.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
normal = doc.styles['Normal']
normal.font.size = Pt(11)
normal.paragraph_format.line_spacing = Pt(15)
normal.paragraph_format.space_after = Pt(7)
doc.styles['Subtitle'].font.italic = False
for name, size in [('Title', 25), ('Subtitle', 12), ('Heading 1', 18), ('Heading 2', 13)]:
    st = doc.styles[name]
    st.font.size = Pt(size)
    st.paragraph_format.space_before = Pt(10)
    st.paragraph_format.space_after = Pt(8)
    st.paragraph_format.keep_with_next = True
doc.styles['Title'].paragraph_format.space_before = Pt(0)
doc.core_properties.title = 'A股雷达模块打磨与独立审计报告'
doc.core_properties.subject = '1.5.2 模块台账与全生命周期审计'
doc.core_properties.author = 'Codex 协作审计'

def p(text, bold_prefix=None, style=None):
    q = doc.add_paragraph(style=style)
    if bold_prefix and text.startswith(bold_prefix):
        q.add_run(bold_prefix).bold = True
        q.add_run(text[len(bold_prefix):])
    else:
        q.add_run(text)
    return q

def h(text):
    doc.add_heading(text, level=1)

def sub(text):
    doc.add_heading(text, level=2)

def page():
    doc.add_page_break()

def loc(file, needle):
    lines = (ROOT / file).read_text().splitlines()
    return f'{file}:{next(i for i, row in enumerate(lines, 1) if needle in row)}'

def issue(code, priority, title, location, trigger, fix, evidence):
    sub(f'{code}  {priority}  {title}')
    q = p('代码位置  ' + location)
    q.paragraph_format.space_after = Pt(5)
    for run in q.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor.from_string('505050')
    p('触发与影响  ' + trigger, '触发与影响  ')
    p('已实施修复  ' + fix, '已实施修复  ')
    p('验证证据  ' + evidence, '验证证据  ')

def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    for i, width in enumerate(widths):
        t.columns[i].width = Inches(width)
    pr = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        b = OxmlElement('w:' + edge)
        b.set(qn('w:val'), 'single'); b.set(qn('w:sz'), '4'); b.set(qn('w:color'), 'D9D9D9')
        borders.append(b)
    pr.append(borders)
    for values in [headers] + rows:
        cells = t.rows[0].cells if values is headers else t.add_row().cells
        for i, (cell, value) in enumerate(zip(cells, values)):
            cell.width = Inches(widths[i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = str(value)
            cp = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for edge in ['top', 'left', 'bottom', 'right']:
                m = OxmlElement('w:' + edge); m.set(qn('w:w'), '90'); m.set(qn('w:type'), 'dxa'); margins.append(m)
            cp.append(margins)
            if values is headers:
                sh = OxmlElement('w:shd'); sh.set(qn('w:fill'), 'DCE7EF'); cp.append(sh)
            for q in cell.paragraphs:
                q.paragraph_format.space_after = Pt(1); q.paragraph_format.line_spacing = Pt(14)
                for run in q.runs:
                    run.font.size = Pt(10)
                    run.bold = values is headers
            nr = OxmlElement('w:cantSplit'); cell._tc.getparent().get_or_add_trPr().append(nr)
    header = OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(header)
    p('')
    return t

def link(label, url):
    q = doc.add_paragraph()
    q.paragraph_format.space_after = Pt(4)
    rel = q.part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    el = OxmlElement('w:hyperlink'); el.set(qn('r:id'), rel)
    run = OxmlElement('w:r'); rp = OxmlElement('w:rPr'); color = OxmlElement('w:color'); color.set(qn('w:val'), '245B82'); rp.append(color)
    size = OxmlElement('w:sz'); size.set(qn('w:val'), '19'); rp.append(size); run.append(rp)
    text = OxmlElement('w:t'); text.text = label; run.append(text); el.append(run); q._p.append(el)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('A股雷达 1.5.2  ·  ')
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
for run in footer.runs:
    run.font.size = Pt(9)

# Every source entry is retained, including overlapping review scopes.
modules = []
for m in fe:
    modules.append({'domain':'frontend', **m, 'source':'frontend-inventory.json'})
for m in be:
    modules.append({'domain':'backend', **m, 'status':'implemented' if m['new_issue'] else 'retained', 'source':'backend-inventory.json'})
for m in rt:
    modules.append({'domain':'root', **m, 'source':'root-inventory.json'})
for m in dep['issues']:
    modules.append({'domain':'dependency', 'module':m['title'], **m, 'source':'dependency-implementation.json'})
issues = []
for m in fe[:5]:
    issues.append({'domain':'frontend', **m})
issues += [{'domain':'backend', **m} for m in impl['backend']['issues']]
issues += [{'domain':'dependency', **m} for m in dep['issues']]
issues += [{'domain':'root', **m} for m in impl['root']['findings']]
ledger = {'version':'1.5.2','status':'draft_pending_final_evidence' if not args.final else 'final_evidence_loaded',
          'sourceEntryCount':len(modules),'domainCounts':dict(Counter(m['domain'] for m in modules)),
          'countingNote':'Source review entries, not unique functional modules. BE-LOG overlaps ROOT-M01; BE-SEARCH has two fixes; independent audit findings are counted separately.',
          'implementationIssueCount':len(issues),'implementationRiskCounts':dict(Counter(m['risk'] for m in issues)),
          'modules':modules,'implementationIssues':issues,
          'independentAudit':audit,'finalValidation':validation,
          'reviewFollowups':[x for x in (audit or {}).get('findings',[]) if x['id'] in (audit or {}).get('verdict',{}).get('newIndependentFindings',[])],
          'limitations':['No formal completeness or zero-vulnerability assurance','Synthetic cache benchmark is not whole-app performance','No real-profile fault injection or long soak','Retained entries describe reviewed boundaries, not proof of every line'],
          'evidenceHashes':{f.name:hashlib.sha256(f.read_bytes()).hexdigest() for f in ART.glob('*-implementation.json')}}
(ART/'report-inventory.json').write_text(json.dumps(ledger,ensure_ascii=False,indent=2)+'\n')

def detailed(code, risk, title, file, needle, trigger, remedy, proof):
    issue(code, risk+'风险', title, loc(file,needle), trigger, remedy, proof)

def retained_rows(items):
    result=[]
    for m in items:
        files=m['files'] if isinstance(m['files'],str) else '、'.join(m['files'])
        reason=m.get('finding') or '；'.join(m.get('verified_protections',[]))
        verify=m.get('verification',m.get('acceptance',''))
        result.append([m['id']+'\n'+m['module']+'\n'+m['risk']+'风险', files, reason+'\n处理：'+m['recommendation']+'\n核验依据：'+verify])
    return result

# 1 — executive scope
p('A股雷达模块打磨与独立审计报告',style='Title')
p('版本 1.5.2    2026年10月1日'+('' if args.final else '    工作草稿'),style='Subtitle')
p('本报告面向软件使用者与维护者，说明从1.5.1出发的小模块检查、兼容修复、局部性能工作，以及未参与施工的独立agent对需求、设计、代码、测试和部署的复核。问题按高、中、低风险列出，检查后保留的模块也逐项说明理由。')
p('本轮共整改18项问题：2项高风险、11项中风险、5项低风险。包括初始14项、独立复核追加的3项交付缺陷及原生启动发现的1项主题读取问题；深层归档和发布入口补修分别归并原问题，避免重复计数。具体风险、整改及验证见后文。')
h('检查范围与计数口径')
table(['领域','来源条目','处理范围'],[
 ['前端','19','5项修复，13项检查保留，1项维护建议'],
 ['后端','10','缓存、搜索、worker三个模块内4项修复；7项保留'],
 ['依赖与门禁','4','安全补丁、内嵌版本核查、发布门禁、安装策略4项整改'],
 ['主进程及交付','3','日志1项、交付3项与建窗1项修复；服务配置读取保留'],
], [1.2,1,4.7])
p('台账共36条来源检查记录，不声称36个独立功能：后端日志与root日志检查存在范围重叠；股票搜索一个模块拆成请求合并和结构校验两个问题。模块风险标签描述本次关注程度，不代表检查后保留者存在同等级未修缺陷。')
sub('风险定义与变更约束')
p('高风险：公开安全公告命中运行时或关键发布控制缺口，优先处理；中风险：可复现的用户流程、数据恢复或资源效率问题；低风险：边界输入、取消原因或维护防护问题。分类用于本项目整改排序，不替代第三方漏洞评级。')
p('不改变策略权重、交易阈值和预设参数；不迁移数据库；保留既有文件格式及1.4.6、1.5.0、1.5.1成品。故障注入使用合成数据、临时目录和隔离浏览器。原生兼容检查与最终归档证据另列，不把计划写成已完成。')

# 2 — traceability
page(); h('全生命周期追溯与职责')
p('审计基准为本轮baseline-source，而不是整个混合工作区的历史未提交diff。初始记录保存250份源码，保留1278个既有成品与配置文件指纹。设计与计划明确采用现有边界内的小步修改。')
table(['阶段','需求与实现对应','证据及判定边界'],[
 ['需求','稳定细小模块；逐项修复及局部优化；独立复核；旧版保留','module-polish-design.md；不扩展为策略收益或全面安全认证'],
 ['设计','renderer/preload/main/service接口不变；纯函数与日志工具提取；缓存保持TTL/LRU/取消语义','module-polish.md；允许测量后不改，避免无依据重写'],
 ['代码','前端5项；后端4项；依赖4项；日志1项；交付3项；建窗1项','四组implementation JSON及相对baseline的diff；修改与保留见逐项台账'],
 ['测试','先失败复现，再定向回归；冻结后统一全量、双入口类型、交互与依赖门禁','实施者日志与test-validation分开，重叠计数不相加'],
 ['独立审计','全新上下文agent复核需求、设计、代码、测试、部署；可执行独立探针','independent_lifecycle_audit未参与被审设计或编码；结论不等同外部机构认证'],
 ['部署','Web、独立复盘、macOS arm64构建，正常原生启动与新目录归档','delivery-native、安装校验及交付验证记录；Windows/Linux实机不预先视为通过'],
], [1,2.7,3.2])
sub('作者与审计者区分')
p('施工协调方负责汇总证据、模块台账和Word排版，不承担此次独立审计。当前独立审计者为independent_lifecycle_audit，接收需求、设计、代码快照及验证证据，在独立上下文中复核。施工方修复新增问题后，由审计方窄范围复验。')
p('最终独立审计结论待证据生成后补入；当前草稿不把实施者自测替代独立结论。' if not audit else '独立审计记录已经生成，其发现、整改与复验闭环见独立审计章节。')

# 3 — front first three
page(); h('前端输入与恢复修复')
for m,needle in zip(fe[:3],['const [composing','const sameStoredShape','function loadArchive']):
    detailed(m['id'],m['risk'],m['module'],m['files'][0],needle,m['finding'],m['recommendation'],m['verification']+'；状态：实施完成，最终冻结验证见后文。')
p('本地JSON与档案恢复保持只读：读取错误主记录时不静默覆盖原始档案。深层校验覆盖个股、市场及决策摘要实际渲染结构；优先保留主档有效行，仅主档无可用记录时恢复备份，兼容合法旧档案。',style=None)

# 4 — front remaining
page(); h('前端数值与上海日期')
for m,needle in zip(fe[3:5],['export function nullableNumber','export function shanghaiDateTagFrom']):
    detailed(m['id'],m['risk'],m['module'],m['files'][0],needle,m['finding'],m['recommendation'],m['verification']+'；状态：实施完成。')
sub('测试证据如何对应真实行为')
p('前端先保留10个预期失败证据：存储和缺失数值5项、搜索和档案4项、日期1项。日期测试执行实际UI保存handler；输入测试覆盖composition期间0次搜索、结束后最终文本1次请求、慢旧请求和卸载失效。')
p('初轮定向测试46项单测、10项浏览器通过；独立审计后的深层修复再通过30项单测、16项浏览器与独立探针及两套类型检查，重叠测试不累计为独立总数。端口占用与DOM换行断言修正属于测试fixture问题，未计为产品缺陷。最终默认全量与最终浏览器结果以统一验证记录为准。')
sub('兼容性边界')
p('nullableNumber只接受有限数字或非空数字字符串，真实零值仍有效；原有数值权重和阈值不变。上海日期统一使用Asia/Shanghai，不再截取UTC日期；无效时间按当前上海日期回退。公共存储键与JSON格式保留，可选validator兼容旧调用。')

# 5/6 — every retained frontend item
page(); h('前端检查后保留模块一')
p('以下模块保留现有行为。列出已检查的保护与对应回归来源，不把保留解释为逐行形式化证明。')
table(['编号与模块','代码位置','检查结果 保留理由与核验'],retained_rows(fe[5:12]),[1.25,1.65,4.0])
page(); h('前端检查后保留模块二')
table(['编号与模块','代码位置','检查结果 保留理由与核验'],retained_rows(fe[12:18]),[1.25,1.65,4.0])
sub('FE M01 中风险维护建议')
p(fe[18]['finding']+'。'+fe[18]['recommendation']+'。状态：建议，未开展广泛组件或CSS重写。')
p('涉及src/App.tsx、src/ProfessionalReview.tsx和src/styles.css。验收先记录首屏、切页、渲染长任务与bundle数据，再按职责提取；单凭文件大小不宣称运行变慢或重构后变快。')

# 7 — backend changes
page(); h('后端缓存 搜索与取消修复')
backend_details=[
 ('BE-CACHE','中','有界缓存热路径','electron/bounded-cache.cjs','  expiryOf(', '5000条新鲜缓存中，100组get与has仍读取100万次无关过期元数据；批量填充出现平方级扫描工作。','维护完成项计数与最早过期点，只在到期或预算边界扫描。保留LRU、字节预算、pending生产者所有权和无常驻定时器。','访问计数红测、最早过期项替换/删除、stale、clear、溢出及pending回归。'),
 ('BE-SEARCH-FLIGHT','中','相同查询请求合并','electron/services.cjs','async function searchSecurities','同时发起同一未缓存查询时，旧实现重复请求上游。','复用serviceRuntime.cached单生产者；维持5分钟TTL、排序、ST过滤和直接代码回退；空结果不缓存。','延迟同词并发一次真实生产者；已完成缓存命中；失败与空结果后可重试。'),
 ('BE-SEARCH-SCHEMA','低','搜索响应结构校验','electron/services.cjs','function normalizeSearchSecurity','Data非数组或结果含null时，TypeError使同批合法建议也无法返回。','显式校验数组，跳过null、数组和非对象行；不扩大证券类型范围。','坏信封、传输失败后成功恢复；混合坏行仍保留合法证券身份。'),
 ('BE-WORKER','低','取消原因不能表示成功','electron/worker-runner.cjs','  function finish(','abort或shutdown原因为null、false、0或空串时，旧代码按真值判断而错误resolve。','以明确error/value结果区分拒绝和成功，保留原因身份；物理exit/terminate前不释放资源。','活动、排队、预取消与shutdown falsy原因均拒绝；drain等待实际结束。')]
for row in backend_details: detailed(*row)

# 8 — retained backend
page(); h('后端检查后保留模块')
table(['编号与模块','代码位置','检查结果 保留理由与核验'],retained_rows(be[3:]),[1.25,1.7,3.95])
p('BE-LOG为后端只读边界检查；root随后负责main日志异常对象处理的修复，二者不是两个重复计数的缺陷。文件主备恢复继续保持上轮方案，不声称跨文件断电绝对原子性。')

# 9 — dependency changes
page(); h('依赖安全与发布门禁整改')
for x,file,needle in zip(dep['issues'],['package.json','qa/verify-dependencies.cjs','qa/verify-dependencies.cjs','pnpm-workspace.yaml'],['"electron"','const undici =','function evaluateAudit','overrides:']):
    title=x['title']; before=x['before']; fix=x['fix']; proof=x['verification']
    translations={
      'DEP-01':('Electron 43.3.0与11个传递版本分支匹配公开公告；dev依赖标签不能排除实际交付运行时。','在同主版本内升级12个精确版本分支，保留直接前端依赖与业务算法。','重新audit及13个直接依赖、12个受影响分支版本核对。'),
      'DEP-02':('首个候选43.5.0虽达到npm审计零匹配，内嵌undici仍为7.29.0。','改用Electron 43.7.5；实际Node24.21.0内嵌undici7.29.1；门禁检查内嵌版本下限。','实际Electron进程process.versions核验，不能用npm图代替二进制证据。'),
      'DEP-03':('原发布只阻断critical；缺少声明、锁、安装和实际二进制一致性检查。','high/critical一律阻断，包含dev包；审计失败、畸形响应、计数矛盾均失败关闭，不设忽略规则。','门禁与发布矩阵30项通过；7个公开完整应用入口在打包前只执行一次审计，失败后停止后续打包。'),
      'DEP-04':('旧nanoid override固定受影响版本；保留了旧Electron发布等待期豁免。','更新精确兼容override，移除旧运行时豁免，保留包龄与registry安全策略。','声明的pnpm11.16.0冻结安装和供应链策略验证成功。')}
    detailed(x['id'],x['risk'],title,file,needle,*translations[x['id']])

# 10 — versions and local measurements
page(); h('实际版本与局部性能证据')
p('依赖实施记录显示：54条公告版本记录、45个唯一GHSA、183条展开路径降至当次审计0匹配。此前34条high、14条moderate、6条low是版本匹配记录，不能表述为54个已利用生产漏洞；当前0匹配也不等于零漏洞。')
table(['依赖','本轮版本变化'],[[x['package'],x['before']+' → '+x['installed']] for x in dep['changes']],[2.2,4.7])
p('实际运行时：Electron 43.7.5，Node 24.21.0，Chromium 150.0.7871.250，内嵌undici 7.29.1；Electron ABI仍为148。34项HTTP与安全测试在新Electron的Node环境执行通过，开发就绪端点和关闭端口deadline探针通过。')
sub('缓存同输入合成基准')
p('macOS arm64、Node24.19.0，5000条稳定过期缓存、20000次命中、5轮：填充中位数130.503ms → 1.496ms；命中中位数1047.331ms → 1.611ms。两版均5000条、校验和49990000；测试证明该合成热路径减少扫描工作。')
p('此基准不是新Electron全应用测速，未测真实行情吞吐、首屏延迟、RSS或长期泄漏，不能外推整机数百倍提速。搜索合并也仅证明相同并发查询减少重复上游工作。')

# 11 — root diagnostics and measured retained
page(); h('主进程诊断与设置读取取舍')
detailed('ROOT-LOG-01','中','故障处理器再次抛错','electron/runtime-diagnostics.cjs','function createRuntimeLogger',impl['root']['findings'][0]['cause'],impl['root']['findings'][0]['fix'],'旧回调5个getter/coercion场景先失败；27项日志、安全及窗口集成契约通过。状态：实施完成。')
sub('ROOT M02 设置读取与建窗主题')
p('使用实际main复盘IPC设置路径，服务替身与合成凭据隔离运行；5轮各500次调用均500次文件读取、1000次解密，耗时中位数25.817ms。重复同步工作已确认，但这里不代表系统Keychain或真实刷新负载耗时。')
p('服务端保留逐次读取的外部文件失效语义；原生启动发现建窗主题提前解密凭据，因此单独修复此路径。只有明确真实负载预算与外部修改契约后才考虑设置快照；后续验收需覆盖保存失败不推进、清密钥立即生效、备份恢复与外部文件修改可见性。没有新增数据库或后台配置缓存。')
detailed('ROOT-STARTUP-01','中','读取窗口主题不应先解密凭据','electron/main.cjs','  // Window appearance needs no credentials.',
         '原生启动在钥匙串授权处等待；createWindow仅为theme调用settingsForService，窗口创建前即解密两个token。',
         '主题直接读取非敏感配置；仅有非空旧密钥待迁移时才检查系统安全存储，空配置或已加密配置的公开读取不初始化钥匙串；真实加解密保护不变。',
         '实际建窗边界先红后绿，30项相关测试通过；独立探针建窗前解密两次降为零；40组迁移结果与基线一致，25组无待迁移秘密组合零系统调用，真实服务读取仍解密且公开返回掩码。')
sub('维护性改善的边界')
p('数值解析提取为纯函数，日志提取为注入环境的诊断模块，搜索加载与缓存订阅分离，worker完成结果改为明确状态。修改围绕已证实问题，不改算法规则。日志仍同步、尽力写入；磁盘故障下不保证日志持久化，也未新增遥测系统。')
sub('官方依据')
for label,url in [('Electron 43.7.5 发行记录',dep['sources'][0]),('Node 24.21.0 发行记录',dep['sources'][1]),('Electron 安全公告',dep['sources'][2]),('xmldom 安全公告',dep['sources'][3]),('undici 安全公告',dep['sources'][4])]: link(label,url)

# Final frozen validation — numbers are read from actual evidence.
page(); h('冻结后的统一验证')
p('前三轮由测试agent独立执行；最后启动补修的第四轮由协调agent运行，独立审计者复核原始日志与代码绑定。多轮结果不重复相加，表内最终命令、轮次和指纹来自test-validation.json。')
if validation and validation.get('status') == 'passed':
    rows=[]
    labels={'full':'默认全量', 'typecheck':'主入口类型', 'typecheck-review':'复盘入口类型', 'ui':'界面回归', 'independent':'独立探针', 'dependencies':'依赖一致性', 'security':'联网安全门禁', 'delivery-contract':'交付边界'}
    for phase in validation.get('finalChecks',[]):
        if phase.get('status') in ('confirmed-defect-awaiting-implementation','failed'):
            continue
        counts=phase.get('counts') or {}
        summary=(str(counts.get('pass',0))+' / '+str(counts.get('tests',0))+' 通过') if counts.get('tests') else '退出码 '+str(phase.get('exitCode','见记录'))
        rows.append([labels.get(phase.get('name'),phase.get('name',''))+('（沿用已验输入）' if phase.get('evidenceReused') else ''),summary,phase.get('log','')])
    table(['验证项','结果','原始证据'],rows,[2.1,1.3,3.5])
    p('混合主档修复后重新执行全量、界面、独立探针和双类型检查；依赖输入未变，测试台账复用首轮依赖证据，构建与安装另执行新联网门禁。第三轮补修建窗主题，第四轮补修迁移可用性检查，均重新运行全量、界面、独立探针、交付契约与双类型检查。')
    p('最终代码清单摘要：'+validation.get('sourceManifestSha256','未提供')[:32]+'…；测试前后核对相同源码。报告生成脚本与文档不属于运行时代码冻结集，成品另以应用完整manifest绑定。')
else:
    p('最终统一测试正在进行，工作草稿不作完成声明。')
sub('独立性能和兼容探针')
p('缓存差分：8个确定性seed，24000步baseline/current操作一致，覆盖TTL、stale、pending、条目/字节上限、LRU及删除清空。它证明所测合成序列保持语义，不代表对所有状态作形式化证明。')
p('归档故障在原始独立3变体探针中先失败，整改后复跑；测试通过还需结合有效旧格式、新格式、市场归档与混合主档优先级案例。依赖门禁验证包含实际命令链中止，而非只检查脚本字符串。')

page(); h('独立审计新增问题与闭环')
if audit:
    ids=set(audit.get('verdict',{}).get('newIndependentFindings',[]))
    for finding in audit.get('findings',[]):
        if finding['id'] not in ids:
            continue
        sub(finding['id']+'  '+finding['risk']+'风险')
        p('问题与影响  '+finding['trigger']+' '+finding['impact'])
        p('整改建议  '+finding['recommendation'])
        p('独立复验  '+finding['fixVerification'])
    p('独立审计结论：代码与隔离资料原生范围已复核，确认代码缺陷已整改；真实凭据的系统授权属于未完成的用户操作。归档校验结论见同目录交付验证记录和最终独立交付复核。' if args.final else '独立审计当前阶段：'+audit.get('stage','详见审计JSON'))
else:
    p('独立审计证据尚未到齐。')

page(); h('交付校验重构与验证边界')
for item in [x for x in impl['root']['findings'] if x['id'].startswith('IA-DELIVERY')]:
    detailed(item['id'],item['risk'],item['title'],'qa/delivery_contract.py',
             {'IA-DELIVERY-BIND-01':'def verify_acceptance','IA-DELIVERY-BIND-02':'def app_manifest','IA-DELIVERY-DATA-01':'def evidence_paths'}[item['id']],
             item['cause'],item['fix'],'交付回归6项通过；独立合成探针分别测试拒绝分支及有效正向基线，复验详情见审计记录。')
sub('未验证与不能外推的范围')
p('本轮未进行全市场实盘扫描、长期soak、Windows/Linux或macOS Intel实机验收。未声称穷尽所有历史数据语义、上游数据源异常或代码状态；没有攻击利用或未知漏洞排除证明。合成缓存提速不代表整应用加速。')
p('已确认缺陷需要整改闭环；后续维护建议与未执行的验证范围单独列出，不能把建议等同遗漏缺陷，也不能把未测项目写成通过。')

# 13 — delivery and follow-on
page(); h('部署保留与后续维护')
sub('ROOT M03 交付保护')
p('新版独立放在桌面a股的“细节打磨版1.5.2”，旧1.4.6、1.5.0、1.5.1成品保留。原生检查按证据中标明的资料目录进行，不保存真实配置、不触发全市场扫描；隔离资料的通过不等同真实凭据授权已完成。1278个保留文件以交付前后指纹核对，不能仅靠目录仍存在认定保护成功。')
if native:
    p(native.get('reportSummary','原生交付记录已生成；最终正文将按其实际平台与检查结果列示。'))
else:
    p('当前草稿的构建、实际包内运行时、安装文件核对、原生启动及保留指纹结果尚待交付证据，不在本页提前标记通过。')
p('本机包未使用Apple Developer ID签名及公证；供本机使用，不视为已达到对外发行条件。主JS约553kB未压缩，保留构建体积提醒；需测量真实首屏后再决定按页面拆包，不凭体积认定运行卡顿。')
sub('维护顺序与验收')
table(['优先级','后续工作','验收方式'],[
 ['高','每次发布重跑严格依赖门禁，检查Electron内嵌库','声明/锁/安装/二进制一致；high/critical阻断；审计服务失败不得忽略'],
 ['中','新增字段持续补齐消费者边界测试','先复现，实施者整改，审计方窄复验；必要时重新冻结相关测试'],
 ['中','大页面与样式按职责分期拆分','先冻结交互契约与指标，单批不混策略改写；不能用行数替代收益指标'],
 ['低','设置快照或进一步缓存优化先测真实工作量','明确预算、外部文件失效与取消所有权，未超预算不无目标引入状态'],
 ['发布前','其他平台原生验收','相应Windows/Linux/mac Intel环境验证后才交付对应成品'],
],[.85,2.35,3.7])
p('源码、证据包与Word报告分别归档；CRC、逐文件SHA-256和包哈希见同目录交付验证记录。报告唯一文档成品为DOCX，页图只用于内部排版检查。证据包排除真实profile和基线源码。')
p('证据索引：report-inventory.json保留全部来源条目；各inventory和implementation记录检查与施工；test-validation记录统一冻结；独立审计记录需求到部署闭环；delivery-native与安装校验记录实际成品。官方依赖依据链接在前页。')

OUT.parent.mkdir(parents=True,exist_ok=True)
for border in list(doc.element.iter(qn('w:pBdr'))): border.getparent().remove(border)
for grid in list(doc.element.iter(qn('w:docGrid'))): grid.getparent().remove(grid)
for paragraph in doc.element.iter(qn('w:p')):
    pr=paragraph.find(qn('w:pPr'))
    if pr is None: pr=OxmlElement('w:pPr'); paragraph.insert(0,pr)
    snap=OxmlElement('w:snapToGrid'); snap.set(qn('w:val'),'0'); pr.append(snap)
doc.save(OUT)
(ART/'word-build.json').write_text(json.dumps({'output':str(OUT),'draft':not args.final,'plannedPages':16,'version':'1.5.2','sourceEntries':len(modules),'implementationIssues':len(issues)},ensure_ascii=False,indent=2)+'\n')
print(OUT)
