"""Build the scoped 1.5.1 audit report from frozen evidence; no production writes."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parents[1]
ART = ROOT / 'artifacts/fullstack-audit-20261001'
OUT = Path('/Users/fengjinwen/Desktop/a股/全栈审计优化版1.5.1/全栈审计与优化报告.docx')
validation = json.loads((ART / 'test-validation.json').read_text())
deps = json.loads((ART / 'dependency-review.json').read_text())
review = json.loads((ART / 'final-review.json').read_text())
native_file = ART / 'delivery-native.json'
native = json.loads(native_file.read_text()) if native_file.exists() else None
assert json.loads((ROOT / 'package.json').read_text())['version'] == '1.5.1'
assert validation['firstFreeze']['default']['pass'] == 702
assert validation['iteration2']['ui']['pass'] == 16
assert review['open_findings'] == []

doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin, section.bottom_margin = Inches(.7), Inches(.65)
section.left_margin = section.right_margin = Inches(.78)
section.footer_distance = Inches(.3)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3']:
    st = doc.styles[name]
    st.font.name = 'Arial Unicode MS'
    st.font.color.rgb = RGBColor(0, 0, 0)
    fonts = st.element.get_or_add_rPr().rFonts
    for attr in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme']:
        fonts.attrib.pop(qn('w:' + attr), None)
    fonts.set(qn('w:eastAsia'), 'Arial Unicode MS')
    for border in list(st.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
normal = doc.styles['Normal']
normal.font.size = Pt(11)
normal.paragraph_format.line_spacing = Pt(15)
normal.paragraph_format.space_after = Pt(7)
doc.styles['Subtitle'].font.italic = False
for name, size in [('Title', 25), ('Subtitle', 12), ('Heading 1', 18), ('Heading 2', 13)]:
    st = doc.styles[name]
    st.font.size = Pt(size)
    st.paragraph_format.space_before = Pt(10)
    st.paragraph_format.space_after = Pt(8)
    st.paragraph_format.keep_with_next = True
doc.styles['Title'].paragraph_format.space_before = Pt(0)
doc.core_properties.title = 'A股雷达全栈审计与优化报告'
doc.core_properties.subject = '1.5.1 兼容性修复与验证范围'
doc.core_properties.author = 'Codex 协作审计'

def p(text, bold_prefix=None, style=None):
    q = doc.add_paragraph(style=style)
    if bold_prefix and text.startswith(bold_prefix):
        q.add_run(bold_prefix).bold = True
        q.add_run(text[len(bold_prefix):])
    else:
        q.add_run(text)
    return q

def h(text):
    doc.add_heading(text, level=1)

def sub(text):
    doc.add_heading(text, level=2)

def page():
    doc.add_page_break()

def loc(file, needle):
    lines = (ROOT / file).read_text().splitlines()
    return f'{file}:{next(i for i, row in enumerate(lines, 1) if needle in row)}'

def issue(code, priority, title, location, trigger, fix, evidence):
    sub(f'{code}  {priority}  {title}')
    q = p('代码位置  ' + location)
    q.paragraph_format.space_after = Pt(5)
    for run in q.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor.from_string('505050')
    p('触发与影响  ' + trigger, '触发与影响  ')
    p('已实施修复  ' + fix, '已实施修复  ')
    p('验证证据  ' + evidence, '验证证据  ')

def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    for i, width in enumerate(widths):
        t.columns[i].width = Inches(width)
    pr = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        b = OxmlElement('w:' + edge)
        b.set(qn('w:val'), 'single'); b.set(qn('w:sz'), '4'); b.set(qn('w:color'), 'D9D9D9')
        borders.append(b)
    pr.append(borders)
    for values in [headers] + rows:
        cells = t.rows[0].cells if values is headers else t.add_row().cells
        for i, (cell, value) in enumerate(zip(cells, values)):
            cell.width = Inches(widths[i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = str(value)
            cp = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for edge in ['top', 'left', 'bottom', 'right']:
                m = OxmlElement('w:' + edge); m.set(qn('w:w'), '90'); m.set(qn('w:type'), 'dxa'); margins.append(m)
            cp.append(margins)
            if values is headers:
                sh = OxmlElement('w:shd'); sh.set(qn('w:fill'), 'DCE7EF'); cp.append(sh)
            for q in cell.paragraphs:
                q.paragraph_format.space_after = Pt(1); q.paragraph_format.line_spacing = Pt(14)
                for run in q.runs:
                    run.font.size = Pt(10)
                    run.bold = values is headers
            nr = OxmlElement('w:cantSplit'); cell._tc.getparent().get_or_add_trPr().append(nr)
    header = OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(header)
    p('')
    return t

def link(label, url):
    q = doc.add_paragraph()
    q.paragraph_format.space_after = Pt(4)
    rel = q.part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    el = OxmlElement('w:hyperlink'); el.set(qn('r:id'), rel)
    run = OxmlElement('w:r'); rp = OxmlElement('w:rPr'); color = OxmlElement('w:color'); color.set(qn('w:val'), '245B82'); rp.append(color)
    size = OxmlElement('w:sz'); size.set(qn('w:val'), '19'); rp.append(size); run.append(rp)
    text = OxmlElement('w:t'); text.text = label; run.append(text); el.append(run); q._p.append(el)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('A股雷达 1.5.1  ·  ')
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
for run in footer.runs:
    run.font.size = Pt(9)

# Page 1
p('A股雷达全栈审计与优化报告', style='Title')
p('交付版本 1.5.1    审计日期 2026年10月1日', style='Subtitle')
p('本报告面向软件使用者与后续维护者，说明本轮确认的问题、已经完成的兼容性修复，以及验证仍有的边界。审查覆盖前端状态、后端行情与任务、文件存储及 Electron 系统边界，目标是避免配置和草稿丢失、纠正数据身份和选股输入口径，并让保存与退出行为可解释。')
p('本轮归并为12项：3项P1、8项P2、1项P3，均已完成对应修复或测试修正。两轮独立复审追加的设置排队覆盖与重复策略组计票边界已回归；最终审查无未解决阻断项。依赖告警和未经测量的架构优化另列后续计划，没有混入本次变更。')
h('范围与处理结果')
table(['领域', '条目', '优先级分布', '结果'], [
    ['前端', 'FE01 至 FE05', '2 P1  3 P2', '初始化 保存 草稿 持仓与票数修复'],
    ['后端与测试', 'BE01 至 BE03', '1 P1  2 P2', '指数身份 队列分类 HTTP断言修正'],
    ['存储', 'DATA01 至 DATA02', '2 P2', '错误形状恢复 敏感设置单次提交'],
    ['架构', 'ARCH01 至 ARCH02', '1 P2  1 P3', '托盘降级 前后端设置契约对齐'],
], [.9, 1.25, 1.3, 3.45])
p('P1表示核心数据或决策输入受损，应优先修复；P2表示可触发的状态、恢复或发布可靠性问题；P3表示较低频的契约一致性问题。此分级用于本项目施工排序，不是第三方漏洞严重性评级。')
sub('系统边界与数据存放')
p('React 渲染层经沙箱 preload 和可信主框架 IPC 调用 Node 服务，网络请求与计算使用有界队列、缓存和 worker。设置、观察池、持仓使用 userData 中的 JSON 与 last-good 备份；计划、监控及部分工作区状态使用 localStorage。项目没有业务数据库或 ORM，Electron 自带运行时组件不等于应用正在使用数据库。')
p('本轮不更改策略评分或预设参数，不引入依赖，不进行数据库迁移。旧桌面1.4.6与1.5.0按保留规则留存，新版单独交付。真实用户配置不作为故障注入样本。')

# Page 2
page(); h('前端状态与编辑保护')
issue('FE01', 'P1', '启动未读完设置时覆盖原配置', loc('src/App.tsx', 'const settingsLoaded ='),
      '设置读取延迟时点击主题或语音快捷动作，原写入器会将默认完整快照写盘，重置刷新参数并清空已保存凭据。',
      '设置读取与持仓、观察池及版本读取解耦；所有保存入口只在设置成功读取后开放。读取失败时阻止默认值写回，并提示重试。',
      '真实 React 隔离浏览器延迟和拒绝 getSettings；确认读取前无写入、原Token和17秒刷新值保留，慢持仓读取不阻塞设置加载。')
issue('FE02', 'P2', '旧确认响应覆盖新设置草稿', loc('src/App.tsx', 'const previousSettings ='),
      '提交设置A后继续输入B，旧响应及外部主题变更会把未保存字段重置，用户的新输入丢失。',
      '区分已确认配置与草稿，外部值只更新未编辑字段。保存成功对仍对应本次提交的字段回填规范化结果，新输入保留。',
      '延迟A响应后输入B，释放响应仍保留B；外部主题变化不覆盖脏字段。未继续编辑时，超范围数值按成功返回值回填。')
issue('FE03', 'P2', '持仓旧任务清空新编辑对象', loc('src/App.tsx', 'const holdingEditRevision ='),
      'A的搜索或保存未结束时改为编辑B，原A完成回调会清空B；同一股票的新输入也受影响。',
      '每次切换或编辑递增修订号，成功回调仅清理未变化的原草稿；更新持仓时在最新已接受集合上执行增改并保留创建时间。',
      '浏览器覆盖A切B及同股续改；独立探针延迟真实saveHoldings，检查B的全部字段保留，随后正常提交仍保存并清空。')

# Page 3
page(); h('前端提交语义与策略票数')
issue('FE04', 'P1', '股票池偏离所选策略与最低票数', loc('src/PortfolioBacktestView.tsx', 'const codesByStrategy ='),
      '选择A和B且要求2票，但响应缺B时原代码降为1票；优化组合含额外策略或不同阈值也被直接复用。',
      '始终保留请求的最低票数；优化组合只有策略集合与票数完全一致才复用。按独立策略ID计票，同股票重复行及重复策略组均只计一次。',
      '缺组、优化超集、阈值不同、重复股票、重复策略组均不得凑票；真实满足A与B两票的样本正常入池。未改变量化评分或业务阈值。')
issue('FE05', 'P2', '保存失败仍改变全局生效设置', loc('src/latestSettingsWriter.ts', 'export function createLatestSettingsWriter'),
      '写盘拒绝后，原乐观值仍成为业务配置，并可能混入后续局部保存。最初审计将产品语义标为待确认，施工明确采用保存成功才提交。',
      '串行写入只在成功响应后更新全局；下一补丁基于最后成功状态合并。失败草稿保留重试，运行仍使用上次确认值。',
      '拒绝写入后离开再返回显示已提交值；后续独立补丁不携带被拒绝字段。保存失败提示与真实生效状态一致。')
sub('两轮复审追加的边界')
p('第一轮审查除了原故障，还核查成功规范化回填、同股续改、持仓最新集合，以及备份准备和主文件提交各阶段失败。测试补足主文件损坏、首次保存失败、回滚再次失败时保留恢复文件与记录降级。')
p('第二轮发现主题保存待确认时，全量表单快照会把旧主题再次写回。现仅提交相对已确认配置的编辑字段，执行时再与最新成功状态合并；显式清空Token仍可提交。独立测试同时补出重复策略组计票问题，现同一策略ID共享去重集合。两项均先保留红测证据，再修复回归。')

# Page 4
page(); h('后端身份边界与请求生命周期')
issue('BE01', 'P1', '专业复盘五个指数被股票校验拒绝', loc('electron/services.cjs', 'async function getReviewIndexChart'),
      '打开专业复盘时，五个指数进入普通股票校验，在发出请求前已失败，随后被转换为空图表，指数价格及均线证据缺失。',
      '新增仅限五个规范市场标识的内部指数图表入口，固定日线与不复权；公共股票、ETF和转债入口仍保持严格校验，不新增指数IPC。',
      '离线集成贯通复盘至真实图表缓存及行情URL，五指数生成历史和MA60；上证000001与深市000001银行隔离，伪造身份继续拒绝，备用源保持身份和不复权。')
issue('BE02', 'P2', '本地队列拥堵被误判为主源宕机', loc('electron/trend-screener-adapter.cjs', 'const localCapacityError ='),
      'QUEUE_FULL或QUEUE_TIMEOUT被改写成数据源不可用，污染当前扫描的failedHosts，后续即使队列恢复也持续跳过主源。',
      '原样传播本地容量错误，不记录源故障、不发起备用请求；备用链遇到同类错误同样保留。已有网络故障和取消策略不变。',
      '两类错误均无备用流量与故障诊断；同一扫描后续请求可重试主源。HTTP授权、限流及取消行为继续受已有测试约束。')
issue('BE03', 'P2', 'HTTP超时测试误判第二请求未结束', loc('electron/http-admission.test.cjs', "test('raw Response body"),
      '两个请求顺序取得响应头，截止时间不同；旧测试用第一个结束后的固定等待推断两者都应关闭。基线676项中该项失败，其余675项通过。',
      '在统一总截止时间内等待两条服务端关闭事件及客户端active归零，不增加固定sleep，不提前消费被忽略的body。',
      '慢响应头180毫秒的红测固定复现时序差；修复后保留ignored.bodyUsed=false再验证读取拒绝，确认覆盖未消费body超时。此项修测试，不宣称修复了网络泄漏。')

# Page 5
page(); h('存储恢复与架构契约')
issue('DATA01', 'P2', '合法JSON错误形状遮蔽备份', loc('electron/main.cjs', 'function isSettingsRecord'),
      'settings.json为null、数组或字符串时语法解析成功，原读取绕过last-good，恢复成默认设置。',
      '读写设置均先校验非空对象且非数组，再执行规范化、合并与凭据迁移。',
      '三类错误形状恢复备份中的17秒和dark，并在后续局部保存时保留原有效字段。')
issue('DATA02', 'P2', '敏感设置两次提交导致结果失真', loc('electron/persistence.cjs', 'if (options.synchronizeBackup === true)'),
      '密钥变更原先连续写两次主文件；第一次已提交但第二次失败，会对调用者返回保存失败。',
      '先准备新敏感状态的备份，再单次rename提交主文件；提交失败恢复旧备份，再次失败保留.rollback.tmp并记录降级。',
      '覆盖备份准备、备份提升、首次与已有主文件提交、损坏主文件及重复回滚故障；主文件只提交一次。多文件保存不等于断电下的绝对事务。')
issue('ARCH01', 'P2', '托盘不可用时隐藏唯一窗口', loc('electron/main.cjs', 'let createdTray;'),
      'Windows托盘构建失败后点关闭，原行为仍隐藏窗口，留下无窗口和托盘入口的后台进程。',
      '只在有效托盘存在时隐藏；失败保留窗口，部分原生托盘分配失败时销毁资源。',
      '执行实际关闭回调，覆盖有效、缺失与已销毁托盘；macOS、Linux及显式退出语义保持。未进行原生Windows运行。')
issue('ARCH02', 'P3', '两层设置规范化语义漂移', loc('src/domain/settings.ts', 'export const normalizeSettings'),
      '非法、null或含空白策略值在主进程与前端产生不同配置。正常完整设置不改变。',
      '以前端契约对齐主进程，统一schema、数值、严格布尔、Token类型与策略trim去重。',
      '真实两份函数跨层fixture覆盖全部默认字段、预设与边界；先5组失败，再8项通过。完整默认止盈3.2保留，缺省输入遵循主进程2.4。')

# Page 6
page(); h('验证范围与证据强度')
p('验证分两个冻结点记录，不把重叠套件简单相加，也不把第一轮全量结果表述为最后一次源码的全量重跑。故障注入使用临时目录、合成数据与隔离浏览器，不读写真实用户配置。')
table(['冻结点', '验证内容', '结果及解释'], [
    ['第一次', '默认全量 702项\n两套TypeScript\n浏览器 14项', '全部通过。涵盖后端、存储、窗口生命周期及原前端修复。'],
    ['第二次', '浏览器 16项\n相关单元 6项\n两套TypeScript\n独立探针 3项', '全部通过。追加排队表单和重复策略组边界；未再次运行默认全量。'],
    ['版本对应', 'SHA 256冻结清单', '第二次只改App.tsx与PortfolioBacktestView.tsx及相关浏览器测试；后端与持久化哈希未变。'],
    ['短时本机回环', '两次各3000毫秒\n各192个请求', '各186成功与6个预期失败，无意外失败；结束后队列、连接及活动资源归零。'],
], [1.0, 1.9, 4.0])
sub('结果如何解释')
p('本机回环峰值客户端租约为8。最终一轮服务端瞬时观察响应峰值为9，涉及断连通知传播，不等同于客户端容量违规。两次事件循环分位数的差异不足以证明普遍性能提升；短时RSS和堆变化也不能证明没有长期泄漏。')
p('独立测试覆盖设置读失败、旧响应、草稿保留、真实持仓保存延迟和票数独立性。安全检查确认可信主框架IPC、沙箱、HTTPS外链、权限默认拒绝及敏感日志脱敏，未把既有保护重复报为本轮问题。')
sub('交付与原生检查')
if native is None:
    p('打包与原生桌面验收尚未形成最终交付证据，本报告不将其标为通过。完成后以delivery-native.json记录的版本、包内源码、启动与退出结果，以及旧版和用户配置完整性核对结果为准。')
else:
    # A parent-provided prose summary is deliberately explicit; do not infer
    # native success merely from the presence of a JSON file.
    p(native.get('reportSummary', '打包与原生验收记录已生成；具体结果见delivery-native.json。本报告尚未将未逐项核对的记录提升为通过结论。'))
    install = json.loads((ART / 'delivery-install.json').read_text())
    p(f"安装校验核对{install['allRegularAppFilesVerified']}个应用常规文件及{install['productionFilesVerified']}个生产文件。Web主入口保留552.82 kB构建警告，已有路由懒加载分块；未据此认定性能回退或完成体积优化。源码和证据包的CRC、逐文件SHA-256及包哈希核验见同目录《交付验证记录.json》。")
sub('未验证范围')
p('没有全市场实盘扫描、长时soak、原生Windows或Linux验收；没有重跑所有历史视觉页面。重复独立I/O故障下可能需要人工使用保留的恢复文件。本轮审查不承诺全面无遗漏、零漏洞或策略收益。')
p('核心证据索引  test-validation.json、validation-freeze-manifest.json、iteration2-freeze-manifest.json、independent-boundary-results.json、final-review.json。均位于 artifacts/fullstack-audit-20261001。')

# Page 7
page(); h('依赖审查与升级候选')
p('依赖清单、锁文件和本机安装的13个直接依赖一致，未发现版本漂移。审计产生54条记录、45个唯一GHSA，涉及183条展开依赖路径；其中34条high、14条moderate、6条low、0条critical。这是受影响版本匹配，不是54个已被利用的生产漏洞。')
p('Electron虽列在devDependencies，实际二进制随应用交付；单看生产依赖审计会漏掉它。Vite、builder和wait-on等主要属于构建或开发链。当前静态审查未证实可利用的生产路径，但这不构成普遍安全证明。本轮没有修改依赖或锁文件。')
table(['候选方案', '建议与可行性', '实施验收'], [
    ['A 运行时补丁\n优先 P1', 'Electron 43.3.0 → 43.5.0，保持同主版本。官方GHSA命中版本已确认，当前远程内容及弹窗限制降低部分可达性。', '隔离更新精确版本和锁；核验包内Electron及Chromium/Node，复测preload IPC 导航 退出和mac打包。'],
    ['B 工具链补丁\n优先 P2', '按现有父范围刷新：xmldom 0.8.15、js-yaml 4.3.2、axios 1.20.0、nanoid 3.3.18等。保持各分支主版本。', '冻结安装后重审告警，验证开发启动 代理 plist与打包；清理过时精确override。'],
    ['C 发布门禁\n优先 P2', '增加清单 锁 安装与实际运行时一致性检查；现有critical阈值不会阻断high记录。', '记录runtime-high的API前提与处置，不以dev=true一概排除。'],
], [1.1, 2.95, 2.85])
p('工具链其余目标：undici 7.29.1及6.28.1、fast-uri 3.1.8、joi 18.2.6、brace-expansion各分支1.1.21、2.1.7、5.0.12。候选已核对父版本范围，但尚未安装或验证二进制兼容；npm层undici更新也不会修补Electron内嵌版本。')
p('Vite保留6.4.3：本次没有直接Vite告警，已核查同主版本候选；没有依据在本轮跨主版本升级。A与B宜分开变更、各自可回退，不在已冻结的1.5.1应用修复里隐式升级。')
sub('官方依据')
link('Electron 43.5.0 官方发行记录', 'https://releases.electronjs.org/release/v43.5.0')
link('Electron GHSA qmv3 fv6v rmhq', 'https://github.com/advisories/GHSA-qmv3-fv6v-rmhq')
link('xmldom GHSA 965w 775f mr7g', 'https://github.com/advisories/GHSA-965w-775f-mr7g')
link('js yaml GHSA 2883 xcg3 v3hh', 'https://github.com/advisories/GHSA-2883-xcg3-v3hh')

# Page 8
page(); h('实施顺序与后续决策')
p('已完成的施工按文件所有权分工：前端负责设置、持仓与组合输入；后端负责指数身份与队列语义；主流程负责文件恢复、敏感提交和托盘；测试独立复现与冻结验证；架构负责跨层契约与最后复审。以本轮baseline-source比较差异，不把历史未提交代码全部归为本次修改。')
table(['顺序与优先级', '范围和投入预估', '可行性与验收'], [
    ['1 已实施', '12项兼容性修复\n两轮边界复审', '保留接口和文件格式。第一冻结全量及第二冻结增量通过；交付验证按第6页证据。'],
    ['2 P1 待实施', 'Electron同主版本补丁\n约1至2天', '方案A独立分支或隔离副本；完整验证与运行时版本核验后再另版交付。'],
    ['3 P2 待实施', '工具链与门禁\n约1至2天', '方案B和C保持兼容范围，告警记录逐项说明。跨平台构建需相应执行环境。'],
    ['4 P3 先测量', '同步配置读取热点\n测量约0.5天', '记录500股票刷新下读取次数 解密耗时与主循环延迟。只有超预算才设计设置快照。'],
    ['5 P3 分期', '大模块职责拆分\n先0.5至1天设计', '从纯设置规则或服务适配器提取；固定契约fixture和回归门禁，每批不混入功能重写。'],
], [1.1, 2.0, 3.8])
sub('架构建议的取舍')
p('settingsForService当前同步读配置并解密，可能在高频刷新或系统密钥库交互时造成主进程延迟，但本轮没有对应硬件负载结论。若测量证实瓶颈，再引入成功提交后更新的主进程快照，明确外部文件修改的失效策略；必须验证保存失败不推进、清密钥立即失效和备份恢复兼容。')
p('主入口552.82kB触发构建警告，但已有业务路由lazy分块。App约9800行、services约7940行及CSS叠加增加维护成本；先测首屏、切页和渲染长任务，再按模块拆分并统一样式层，不凭体积宣称速度提升。数据库仅在事务、查询或规模需求明确后另行评估。')
sub('交付保留与追溯')
p('新版目录为桌面a股下“全栈审计优化版1.5.1”，旧1.4.6及“重构版1.5.0”不覆盖。完整性以交付前后哈希核对为准；应用源码、验收证据与报告应一起归档。用户凭据不进入报告，测试仅使用合成标记。')
p('归档文件“审计验证证据.zip”，收录选定JSON、日志与diff，排除真实profile和基线源码。阅读顺序：审计 → implementation → test-validation → final-review → delivery-native；依赖见dependency-review。候选方案不视为已部署。')
link('Electron 安全开发官方建议', 'https://www.electronjs.org/docs/latest/tutorial/security')
link('Electron safeStorage 官方文档', 'https://www.electronjs.org/docs/latest/api/safe-storage')
link('Vite npm 官方版本记录', 'https://registry.npmjs.org/vite')

OUT.parent.mkdir(parents=True, exist_ok=True)
for border in list(doc.element.iter(qn('w:pBdr'))):
    border.getparent().remove(border)
for grid in list(doc.element.iter(qn('w:docGrid'))):
    grid.getparent().remove(grid)
for paragraph in doc.element.iter(qn('w:p')):
    pr = paragraph.find(qn('w:pPr'))
    if pr is None:
        pr = OxmlElement('w:pPr'); paragraph.insert(0, pr)
    snap = OxmlElement('w:snapToGrid'); snap.set(qn('w:val'), '0'); pr.append(snap)
doc.save(OUT)
(ART / 'word-build.json').write_text(json.dumps({'output':str(OUT), 'plannedPages':8, 'nativeEvidenceIncluded':native is not None, 'version':'1.5.1'}, ensure_ascii=False, indent=2)+'\n')
print(OUT)
