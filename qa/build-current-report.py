from pathlib import Path
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json, re
ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'deliverables/current/A股雷达稳定性审核与优化报告.docx'
MD = ROOT / 'docs/current-delivery-report.md'
qa = ROOT / 'artifacts/sector-classification-v144'
log = (qa / 'final-tests.tap').read_text()
assert '# fail 0' in log
count = int(re.search(r'# pass (\d+)', log).group(1))
uilog = (qa / 'independent-ui-regression.tap').read_text()
assert '# fail 0' in uilog
uicount = int(re.search(r'# pass (\d+)', uilog).group(1))
soakpath = qa / 'soak-8h/summary.json'
if not soakpath.exists(): soakpath = qa / 'soak-8h/checkpoint.json'
soak = json.loads(soakpath.read_text()) if soakpath.exists() else {}
soaktext = (f"本版八小时离线连续验证完成，共 {soak.get('completedWaves')} 轮，状态 {soak.get('status')}。" if soak.get('eightHoursCompleted') and soak.get('status') == 'passed' else '本版尚无八小时连续运行通过结论。上一版长测因本次功能更新主动停止；本版源码冻结后重新完整计时，具体状态见验收记录中的 soak-8h。')
pages = [
[
('title','A股雷达稳定性审核与优化报告'),
('subtitle','桌面交付版本 1 4 4'),
('p','本版修正板块分类来源，支持输入中文股票名称查行业，并调整深浅主题的文字和背景。当前可读取同花顺公开行业及概念目录，查询指定股票的官方行业路径；完整一至三级行业树仍未取得，不能把公开目录数量理解为全量三级分类。'),
('h1','1 板块数据修正'),
('p','删除原先由六只示例股票生成的默认分类目录，以及人工编排的电子布、电子树脂和光刻胶材料入口。默认显示全部可读取的同花顺公开板块；搜索框可以按行业或概念名称筛选。'),
('p','本次真实网络核验读取到90个行业入口、361个概念入口。这些是公开页面链接，来源没有标明行业层级。程序按实际响应动态解析，不写死数量，也不把两个平面目录拼成完整行业树。'),
('p','输入股票名称后选择候选，再点击“查行业”。页面从该股票同花顺F10字段读取一级、二级、三级路径，自动打开最深层分类。原六位代码输入方式继续可用，中文输入法组词期间不会提前提交。'),
('h2','完整分类的数据条件'),
('p','同花顺SuperMind官方提供get_industry_relate行业枚举能力，但属于其授权研究环境，不能直接作为现有iFinD HTTP接口调用。本机未配置可用数据权限。公开全股票页面的尾页请求返回HTTP 401，因此目前无法核验全市场名册并逐股补齐。'),
('p','全量接入的下一步需要确认可用的iFinD或SuperMind权限，再核对行业ID、层级、日期及成分覆盖。程序保留来源不可用和部分资料状态；不会用旧式行业列表、人工分组或少量代表股票填补缺口。'),
('h2','数据口径'),
('p','F10同行财务对比表显示报告期和已取得股票，不当作当前完整指数成分。板块页缺少明确行情日期时显示“行情日期未提供”。分页失败保留已经读取的数据，标注不完整，不用样本估算整个板块强度。'),
],
[
('h1','2 界面与操作'),
('p','板块列表保留名称和选中状态；打开板块后主要显示标题、行情、股票表格。原来常驻的长说明、小字注释和来源解释收入“来源与覆盖”折叠区。数据缺失、部分资料和日期仍可直接看到。'),
('p','移除远程Google字体依赖，优先使用macOS系统字体及苹方，Windows使用系统中文字体。正文与主要表格放大到15像素，常用控件和次要文字提高到13至14像素。深浅主题均采用实色表面、更清楚的边框与文字，取消顶栏背景模糊。保留用户已经选择的主题。'),
('h2','查询与请求隔离'),
('p','名称搜索有250毫秒防抖，支持方向键选择、Enter确认、Escape收起。修改名称、切换板块和取消操作会作废旧查询意图；晚到结果不能覆盖后来的选择。ST及非普通A股在候选和成分名单中按现有规则剔除。'),
('h1','3 稳定性保护'),
('p','板块服务保留最多两个实际网络请求，同时限制逻辑请求数量为32，总处理期限为180秒。取消后等待底层请求实际结束才归还并发容量；退出应用会停止并等待未完成请求。'),
('p','新增板块缓存128项和8 MiB序列化预算，目录登记4096项和2 MiB预算；请求结束清理代次登记。强制刷新乱序返回时旧代不能覆盖新代，过大的缓存结果不会长期驻留。字节预算衡量缓存快照，不是整个应用内存上限。'),
('p','既有任务订阅隔离、主进程所有权检查、worker退出清理、共享缓存取消、原生退出重入修复继续保留。用户观察池、持仓、计划及凭据目录不随应用替换而删除。'),
('h2','多代理复审'),
('p','后端代理核验来源、解析及缓存生命周期；前端代理实现名称查询和页面精简并构造回归；独立验收代理测试浏览器交互、两种主题和三种宽度。主代理复核各项修改、运行整体套件并校验安装文件。'),
],
[
('h1','4 当前验证与交付'),
('p',f'当前整体回归{count}项通过，既有真实浏览器稳定性回归{uicount}项通过。板块专项测试另验证目录、股票名称选择、代码查询、输入法、晚到结果隔离、来源折叠及深浅主题显示。专项用例与整体套件存在重合，不重复相加。'),
('p','界面验证覆盖760、1024、1480像素宽度，检查页面横向溢出、表格滚动、实际字号和文字对比度。官方页面抓取用于核对数据契约；浏览器回放和合成故障用于验证交互及错误状态，不代表持续实盘行情可用性。'),
('p',soaktext),
('h2','交付位置与保留规则'),
('p','桌面a股文件夹保留当前应用、源码ZIP、本报告、策略配置及当前验收记录。新版本通过验证后直接替换旧成品，不创建旧版备份。应用中的长期用户数据独立保留。'),
('p','当前源码文件为electron/sector-explorer.cjs、src/SectorExplorerPanel.tsx和src/sector-explorer.css；全局文字与配色位于styles.css、terminal-ui.css、terminal-research.css及workbench.css。验收记录保留网络探针、测试日志、截图、构建与安装校验。'),
('h2','官方接口依据'),
('p','同花顺公开行业目录 https://q.10jqka.com.cn/thshy/'),
('p','同花顺公开概念目录 https://q.10jqka.com.cn/gn/'),
('p','SuperMind API文档 https://quant.10jqka.com.cn/view/help/4'),
('p','iFinD HTTP帮助 https://quantapi.51ifind.com/gwstatic/static/ds_web/quantapi-web/help-center/manual.html'),
('h2','验证范围'),
('p','本机原生验证环境为macOS arm64。未取得Windows及Linux原生运行证据，亦未验证强制断电或多小时真实供应商可用性。测试通过说明已覆盖用例表现正确，不意味着不存在未知缺陷。'),
]
]
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Mm(210), Mm(297)
sec.top_margin, sec.bottom_margin = Mm(18), Mm(18)
sec.left_margin, sec.right_margin = Mm(22), Mm(22)
sec.footer_distance = Mm(9)
styles = doc.styles
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2']:
    s=styles[name]
    s.font.name='Arial'
    s.font.color.rgb=RGBColor(0,0,0)
    rf=s.element.get_or_add_rPr().get_or_add_rFonts()
    for key in ['eastAsiaTheme','asciiTheme','hAnsiTheme','cstheme']:
        if qn('w:'+key) in rf.attrib: del rf.attrib[qn('w:'+key)]
    rf.set(qn('w:eastAsia'),'Songti SC' if name=='Normal' else 'Heiti SC')
    s.paragraph_format.space_after=Pt(6)
    s.paragraph_format.line_spacing=1.10
styles['Normal'].font.size=Pt(10.5)
styles['Title'].font.size=Pt(25)
styles['Title'].font.bold=True
styles['Title'].paragraph_format.space_after=Pt(9)
styles['Subtitle'].font.size=Pt(10)
styles['Subtitle'].paragraph_format.space_after=Pt(15)
styles['Heading 1'].font.size=Pt(16)
styles['Heading 1'].font.bold=True
styles['Heading 1'].paragraph_format.space_before=Pt(9)
styles['Heading 1'].paragraph_format.space_after=Pt(9)
styles['Heading 2'].font.size=Pt(12)
styles['Heading 2'].font.bold=True
styles['Heading 2'].paragraph_format.space_before=Pt(7)
styles['Heading 2'].paragraph_format.space_after=Pt(6)
for name in ['Title','Subtitle','Heading 1','Heading 2']:
    styles[name].paragraph_format.keep_with_next=True
    ppr=styles[name].element.get_or_add_pPr()
    for border in list(ppr.findall(qn('w:pBdr'))): ppr.remove(border)
styles['Normal'].paragraph_format.widow_control=True
doc.core_properties.title='A股雷达稳定性审核与优化报告'
doc.core_properties.subject='1.4.4板块数据与界面更新验收'
doc.core_properties.author=''
doc.core_properties.keywords='A股,稳定性,并发,恢复,验收'
footer=sec.footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=footer.add_run('第 '); r.font.size=Pt(9)
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); footer._p.append(fld)
r=footer.add_run(' 页');r.font.size=Pt(9)

markdown=[]
for pi,blocks in enumerate(pages):
    if pi:
        markdown.append('\n<!-- pagebreak -->\n')
    pending_break = pi > 0
    for kind,content in blocks:
        if kind in ['title','subtitle','h1','h2','p','code']:
            style={'title':'Title','subtitle':'Subtitle','h1':'Heading 1','h2':'Heading 2','p':'Normal','code':'Normal'}[kind]
            p=doc.add_paragraph(content,style=style)
            if pending_break:
                p.paragraph_format.page_break_before=True
                pending_break=False
            if kind=='code':
                p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(9)
                for r in p.runs:r.font.name='Menlo';r.font.size=Pt(9)
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'F2F5F7');p._p.get_or_add_pPr().append(shade)
                markdown.append('```js\n'+content+'\n```\n')
            else:
                prefix={'title':'# ','subtitle':'','h1':'## ','h2':'### ','p':''}[kind]
                markdown.append(prefix+content+'\n')
        else:
            headers,rows,widths=content
            table=doc.add_table(rows=1,cols=len(headers))
            table.alignment=WD_TABLE_ALIGNMENT.CENTER
            table.autofit=False
            for i,w in enumerate(widths):table.columns[i].width=Mm(w)
            for cells,values,ishead in [(table.rows[0].cells,headers,True)]+[(table.add_row().cells,row,False) for row in rows]:
                for ci,(cell,value) in enumerate(zip(cells,values)):
                    cell.width=Mm(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    cell.text=value
                    tcpr=cell._tc.get_or_add_tcPr()
                    margins=OxmlElement('w:tcMar')
                    for side in ['top','bottom','left','right']:
                        e=OxmlElement('w:'+side);e.set(qn('w:w'),'100');e.set(qn('w:type'),'dxa');margins.append(e)
                    tcpr.append(margins)
                    borders=OxmlElement('w:tcBorders')
                    for side in ['top','bottom','left','right']:
                        e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
                    tcpr.append(borders)
                    if ishead:
                        shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DCE6EE');tcpr.append(shade)
                    for p in cell.paragraphs:
                        p.paragraph_format.space_after=Pt(0)
                        p.paragraph_format.line_spacing=1.10
                        p.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 or len(headers)==2 else WD_ALIGN_PARAGRAPH.CENTER
                        for r in p.runs:r.font.size=Pt(9.5);r.font.bold=ishead
            rowpr=table.rows[0]._tr.get_or_add_trPr();repeat=OxmlElement('w:tblHeader');rowpr.append(repeat)
            for row in table.rows:
                e=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(e)
            p=doc.add_paragraph();p.paragraph_format.space_after=Pt(1);p.paragraph_format.line_spacing=0.2;p.add_run().font.size=Pt(2)
            markdown.append('| '+' | '.join(headers)+' |\n| '+' | '.join(['---']*len(headers))+' |\n'+'\n'.join('| '+' | '.join(x.replace('\n','<br>') for x in row)+' |' for row in rows)+'\n')
OUT.parent.mkdir(parents=True,exist_ok=True)
doc.save(OUT)
MD.write_text('\n'.join(markdown),encoding='utf-8')
print(OUT)
print(MD)
